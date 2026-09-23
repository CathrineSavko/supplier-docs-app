import argparse
import ast
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_BUYER = [
    'OOO "SADOVAYA TEKHNIKA I INSTRUMENTY"',
    'INNER CITY TER., BASMANNY MUNICIPAL DISTRICT,',
    'BOLSHAYA POCHTOVAYA STR., 40, BUILDING 1',
    '105082 MOSCOW, RUSSIA, INN 7718215870 KPP 770101001',
]
EXPECTED_AGREEMENT = '№101-L от 06.06.2024'
EXPECTED_PAYMENT = '100% in 25 calendar days after the shipment'


def update_progress(job_dir, percent, stage):
    request = job_dir / 'request.json'
    try:
        data = json.loads(request.read_text(encoding='utf-8'))
        data['progress'] = percent
        data['stage'] = stage
        request.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    except Exception:
        pass


def text(value):
    return '' if value is None else str(value).strip()


def normalized(value):
    return re.sub(r'\s+', ' ', text(value).replace('\xa0', ' ')).strip().casefold()


def compact_article(value):
    return re.sub(r'[^a-z0-9]', '', text(value).casefold())


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def cell_location(cell):
    return f'{cell.coordinate}'


def finding(items, point, sheet, location, found, proposal, correction=None):
    items.append({
        'point': point, 'sheet': sheet, 'location': location,
        'found': found, 'proposal': proposal, 'correction': correction,
    })


def find_cell(ws, predicate, start_row=1):
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if predicate(normalized(cell.value)):
                return cell
    return None


def find_header(ws, words):
    return find_cell(ws, lambda value: all(word in value for word in words))


def table(ws):
    description = find_header(ws, ['description'])
    item = find_header(ws, ['item', 'no'])
    if not description or not item:
        return None
    header_row = max(description.row, item.row)
    total_row = None
    for row in range(header_row + 1, ws.max_row + 1):
        values = [normalized(ws.cell(row, column).value) for column in range(1, ws.max_column + 1)]
        if any(value == 'total' or value.startswith('total ') or value.startswith('total:') for value in values):
            total_row = row
            break
    if total_row is None:
        return None
    rows = [row for row in range(header_row + 1, total_row) if text(ws.cell(row, item.column).value)]
    return {
        'header_row': header_row, 'total_row': total_row, 'rows': rows,
        'item_col': item.column, 'description_col': description.column,
    }


def value_right_of(ws, label):
    for column in range(label.column + 1, ws.max_column + 1):
        candidate = ws.cell(label.row, column)
        if text(candidate.value):
            return candidate
    for row in range(label.row + 1, min(label.row + 3, ws.max_row + 1)):
        candidate = ws.cell(row, label.column)
        if text(candidate.value):
            return candidate
    return None


def buyer_lines(ws):
    label = find_cell(ws, lambda value: 'buyer' in value or 'consignee' in value)
    if not label:
        return None, []
    lines = []
    for row in range(label.row + 1, min(label.row + 7, ws.max_row + 1)):
        value = text(ws.cell(row, label.column).value)
        if value:
            lines.append((value, row))
        if len(lines) == 4:
            break
    return label, lines


def safe_formula_value(ws, cell, visited=None):
    """Evaluate only arithmetic cell formulas used in current supplier tables."""
    value = cell.value
    if is_number(value):
        return float(value)
    if not isinstance(value, str) or not value.startswith('='):
        return None
    visited = visited or set()
    if cell.coordinate in visited:
        return None
    visited.add(cell.coordinate)
    formula = value[1:].upper().replace('$', '')

    def replace_sum(match):
        start, end = match.group(1), match.group(2)
        try:
            start_cell = ws[start]
            end_cell = ws[end]
            result = 0.0
            for row in ws.iter_rows(min_row=start_cell.row, max_row=end_cell.row, min_col=start_cell.column, max_col=end_cell.column):
                for part in row:
                    number = safe_formula_value(ws, part, visited.copy())
                    if number is None:
                        return 'INVALID'
                    result += number
            return str(result)
        except Exception:
            return 'INVALID'

    formula = re.sub(r'SUM\(([A-Z]+\d+):([A-Z]+\d+)\)', replace_sum, formula)

    def replace_ref(match):
        ref = match.group(0)
        if '!' in ref:
            return 'INVALID'
        try:
            number = safe_formula_value(ws, ws[ref], visited.copy())
            return 'INVALID' if number is None else str(number)
        except Exception:
            return 'INVALID'

    formula = re.sub(r'(?<![A-Z!])([A-Z]+\d+)', replace_ref, formula)
    if 'INVALID' in formula or not re.fullmatch(r'[0-9. +\-*/()]+', formula):
        return None
    try:
        node = ast.parse(formula, mode='eval')
        if not all(isinstance(part, (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Add, ast.Sub, ast.Mult, ast.Div, ast.USub, ast.Constant)) for part in ast.walk(node)):
            return None
        result = eval(compile(node, '<formula>', 'eval'), {'__builtins__': {}}, {})
        return float(result) if is_number(result) else None
    except Exception:
        return None


def check_description_text(items, ws, context):
    for row in context['rows']:
        description = ws.cell(row, context['description_col'])
        item = ws.cell(row, context['item_col'])
        value = text(description.value)
        if not value:
            finding(items, 1, ws.title, cell_location(description), 'Поле Description не заполнено.', 'Добавить наименование товара на английском языке.')
            continue
        if re.search(r' {2,}|\s+[,.]', value) or re.search(r'\b(\w+)\s+\1\b', value, flags=re.I):
            finding(items, 1, ws.title, cell_location(description), f'Описание: {value}', 'Убрать повторяющиеся пробелы, слова и пробелы перед знаками препинания.', {'type': 'collapse_spaces', 'cells': [cell_location(description)]})
        article = compact_article(item.value)
        if article and article in compact_article(value):
            finding(items, 10, ws.title, cell_location(description), f'В Description указан артикул {text(item.value)}: {value}', 'Удалить артикул из Description, сохранив наименование товара.')


def check_document(ws, items):
    context = table(ws)
    if not context:
        finding(items, '1–11', ws.title, 'table', 'The item table or Total row was not recognised.', 'Check the standard headers Item No., Description and Total manually.')
        return None
    check_description_text(items, ws, context)
    number_col = context['item_col'] - 1
    if number_col >= 1:
        for expected, row in enumerate(context['rows'], start=1):
            cell = ws.cell(row, number_col)
            if str(cell.value).strip() != str(expected):
                finding(items, 6, ws.title, cell_location(cell), f'Item line number is {text(cell.value)}.', f'Set the line number to {expected}.')
    return context


def check_invoice(ws, context, items):
    trademark = find_header(ws, ['trademark'])
    if trademark:
        for row in context['rows']:
            description = normalized(ws.cell(row, context['description_col']).value)
            mark = text(ws.cell(row, trademark.column).value)
            for brand in ('URZUS', 'KEIZEN'):
                if brand.casefold() in description and brand.casefold() not in normalized(mark):
                    finding(items, 2, ws.title, f'{ws.cell(row, context["description_col"]).coordinate}, {ws.cell(row, trademark.column).coordinate}', f'Description contains {brand}, Trademark is {mark or "empty"}.', f'Set Trademark to {brand}.')
    else:
        finding(items, 2, ws.title, 'Trademark', 'Trademark column was not recognised.', 'Check URZUS/KEIZEN trademarks manually.')

    for point, label_text, required in [(3, 'agreement', EXPECTED_AGREEMENT), (4, 'payment terms', EXPECTED_PAYMENT)]:
        label = find_cell(ws, lambda value, needle=label_text: needle in value)
        value = value_right_of(ws, label) if label else None
        if not value or normalized(value.value) != normalized(required):
            location = cell_location(value) if value else (cell_location(label) if label else label_text)
            found = text(value.value) if value else 'Value is missing.'
            finding(items, point, ws.title, location, f'Current value: {found}', f'Set the value to: {required}')


def check_packing_list(ws, context, items):
    cbm = find_header(ws, ['cbm'])
    if cbm:
        for row in context['rows']:
            cell = ws.cell(row, cbm.column)
            if isinstance(cell.value, str) and cell.value.startswith('='):
                value = safe_formula_value(ws, cell)
                finding(items, 7, ws.title, cell_location(cell), f'В ячейке CBM установлена формула {cell.value}.', f'Заменить формулу рассчитанным значением {value:g}.', {'type': 'formula_to_value', 'cells': [cell_location(cell)], 'value': value})
    net = find_header(ws, ['net', 'weight', 'total'])
    gross = find_header(ws, ['gross', 'weight', 'total'])
    if net and gross:
        for row in context['rows']:
            net_cell, gross_cell = ws.cell(row, net.column), ws.cell(row, gross.column)
            net_value, gross_value = safe_formula_value(ws, net_cell), safe_formula_value(ws, gross_cell)
            if net_value is not None and gross_value is not None and abs(net_value - gross_value) > 0.00001:
                finding(items, 8, ws.title, f'{net_cell.coordinate}, {gross_cell.coordinate}', f'Net Weight Total is {net_value:g}; Gross Weight Total is {gross_value:g}.', 'Confirm which value is correct, then set both fields to that exact value.')
    else:
        finding(items, 8, ws.title, 'Net/Gross Weight Total', 'One or both weight columns were not recognised.', 'Check Net Weight Total and Gross Weight Total manually.')


def check_totals(ws, context, items):
    for column in range(1, ws.max_column + 1):
        values = [safe_formula_value(ws, ws.cell(row, column)) for row in context['rows']]
        numeric = [value for value in values if value is not None]
        if not numeric:
            continue
        total_cell = ws.cell(context['total_row'], column)
        total_value = safe_formula_value(ws, total_cell)
        if total_value is None:
            continue
        expected = sum(numeric)
        if abs(total_value - expected) > 0.00001:
            finding(items, 9, ws.title, cell_location(total_cell), f'Total is {total_value:g}; item lines sum to {expected:g}.', f'Set Total to {expected:g} or correct the item lines.')


def inspect_file(path: Path):
    items = []
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        required = ['INV', 'PL']
        for name in required:
            if name not in workbook.sheetnames:
                finding(items, '1–11', name, 'sheet', f'Sheet {name} is missing.', f'Add or upload a workbook containing sheet {name}.')
        contexts = {}
        for name in required:
            if name not in workbook.sheetnames:
                continue
            ws = workbook[name]
            buyer_label, buyer = buyer_lines(ws)
            actual = [line for line, _ in buyer]
            location = ', '.join(f'A{row}' for _, row in buyer) if buyer else (cell_location(buyer_label) if buyer_label else 'Buyer / Consignee')
            if len(actual) != 4 or [normalized(value) for value in actual] != [normalized(value) for value in EXPECTED_BUYER]:
                finding(items, 5, name, location, 'Buyer / Consignee lines do not exactly match the required text.', 'Use exactly: ' + ' | '.join(EXPECTED_BUYER))
            if any(re.search(r' {2,}', line) for line in actual):
                cells = [f'{buyer_label.column_letter}{row}' for _, row in buyer] if buyer_label else []
                finding(items, 5, name, location, 'В строках Buyer / Consignee есть двойные пробелы.', 'Убрать лишние пробелы внутри строк.', {'type': 'collapse_spaces', 'cells': cells})
            context = check_document(ws, items)
            contexts[name] = context
            if context:
                check_totals(ws, context, items)
        if contexts.get('INV'):
            check_invoice(workbook['INV'], contexts['INV'], items)
        if contexts.get('PL'):
            check_packing_list(workbook['PL'], contexts['PL'], items)
        if contexts.get('INV') and contexts.get('PL'):
            inv, pl = contexts['INV'], contexts['PL']
            inv_rows = {text(workbook['INV'].cell(row, inv['item_col']).value): text(workbook['INV'].cell(row, inv['description_col']).value) for row in inv['rows']}
            pl_rows = {text(workbook['PL'].cell(row, pl['item_col']).value): (row, text(workbook['PL'].cell(row, pl['description_col']).value)) for row in pl['rows']}
            for item, inv_description in inv_rows.items():
                if item in pl_rows and normalized(inv_description) != normalized(pl_rows[item][1]):
                    finding(items, 11, 'INV / PL', f'{item}: INV and PL', f'INV: {inv_description}; PL: {pl_rows[item][1]}', f'Use one identical Description for Item No. {item}.')
    finally:
        workbook.close()
    return {'file': path.name, 'findings': items, 'grammarNote': 'Локальная проверка грамматики выявляет очевидные ошибки: повторяющиеся пробелы, слова и знаки препинания. Описания без замечаний всё равно требуют языковой проверки перед изменением.'}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--job-dir', required=True)
    parser.add_argument('--complete-job', action='store_true')
    args = parser.parse_args()
    job_dir = Path(args.job_dir)
    sources = list((job_dir / 'input').glob('*.xlsx'))
    if not sources:
        raise ValueError('Task 1 needs source .xlsx files.')
    report = []
    for index, source in enumerate(sources, start=1):
        update_progress(job_dir, 5 + round((index - 1) / len(sources) * 90), f'Проверяю файл {index} из {len(sources)}: {source.name}')
        report.append(inspect_file(source))
        update_progress(job_dir, 5 + round(index / len(sources) * 90), f'Проверено файлов: {index} из {len(sources)}')
    result = {'report': report, 'notes': ['Read-only check completed. Source files were not changed.']}
    if args.complete_job:
        request = job_dir / 'request.json'
        data = json.loads(request.read_text(encoding='utf-8'))
        data.update({'status': 'completed', 'progress': 100, 'stage': 'Готово', 'report': report, 'outputFiles': [], 'errors': [], 'notes': result['notes']})
        request.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=True))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)

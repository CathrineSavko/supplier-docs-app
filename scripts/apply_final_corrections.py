import argparse, json, re, shutil
from pathlib import Path
from openpyxl import load_workbook

def main():
    p = argparse.ArgumentParser(); p.add_argument('--job-dir', required=True); args = p.parse_args()
    job_dir = Path(args.job_dir); request = json.loads((job_dir / 'request.json').read_text(encoding='utf-8'))
    selected = request.get('parameters', {}).get('selected', [])
    if not selected: raise ValueError('Не выбраны исправления.')
    grouped = {}
    for issue in selected: grouped.setdefault(issue['file'], []).append(issue)
    output_dir = job_dir / 'output'; output_dir.mkdir(exist_ok=True)
    files, notes = [], []
    for name, issues in grouped.items():
        source = job_dir / 'input' / name; destination = output_dir / name
        shutil.copyfile(source, destination)
        wb = load_workbook(destination, data_only=False)
        try:
            for issue in issues:
                ws = wb[issue['sheet']]; correction = issue['correction']
                for cell_ref in correction['cells']:
                    cell = ws[cell_ref]
                    if correction['type'] == 'collapse_spaces': cell.value = re.sub(r' {2,}', ' ', str(cell.value))
                    elif correction['type'] == 'formula_to_value': cell.value = correction['value']
            wb.save(destination)
        finally: wb.close()
        files.append(name); notes.append(f'{name}: применены выбранные исправления.')
    print(json.dumps({'outputFiles': files, 'notes': notes}, ensure_ascii=True))

if __name__ == '__main__':
    try: main()
    except Exception as error: print(str(error)); raise

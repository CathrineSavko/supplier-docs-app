import argparse
import json
import re
import shutil
import sys
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


def update_progress(job_dir: Path, percent: int, stage: str):
    request = job_dir / 'request.json'
    try:
        data = json.loads(request.read_text(encoding='utf-8'))
        data['progress'] = percent
        data['stage'] = stage
        request.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    except Exception:
        pass


def output_name(stem: str) -> str:
    if re.match(r"^\d+\s*[A-Za-z]+", stem):
        return f"{stem}.xlsx"
    number = re.search(r"\d+", stem)
    letters = re.search(r"[A-Za-z]+", stem)
    if number and letters:
        return f"{number.group(0)} {letters.group(0).upper()}.xlsx"
    return f"{stem}.xlsx"


def find_text(ws, needle: str):
    normalized = re.sub(r"\s+", " ", needle).lower()
    for row in range(1, 34):
        for column in range(1, 10):
            value = ws.cell(row, column).value
            if value is not None and normalized in re.sub(r"\s+", " ", str(value)).lower():
                return ws.cell(row, column)
    return None


def clear_outside_area(ws):
    for cell in list(ws._cells.values()):
        if cell.row > 33 or cell.column > 9:
            cell.value = None
    ws._images = []


def unmerge_intersecting(ws, min_row, max_row, min_col, max_col):
    for merged in list(ws.merged_cells.ranges):
        if not (merged.max_row < min_row or merged.min_row > max_row or merged.max_col < min_col or merged.min_col > max_col):
            ws.unmerge_cells(str(merged))


def copy_block(template_ws, ws, start_row):
    source_first, source_last = 21, 28
    unmerge_intersecting(ws, start_row, start_row + 7, 1, 8)
    for source_row in range(source_first, source_last + 1):
        target_row = start_row + source_row - source_first
        ws.row_dimensions[target_row].height = template_ws.row_dimensions[source_row].height
        for column in range(1, 9):
            source = template_ws.cell(source_row, column)
            target = ws.cell(target_row, column)
            target.value = source.value
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            target.alignment = copy(source.alignment)
            target.protection = copy(source.protection)
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
    for merged in template_ws.merged_cells.ranges:
        if merged.min_row >= source_first and merged.max_row <= source_last and merged.min_col >= 1 and merged.max_col <= 8:
            ws.merge_cells(
                start_row=start_row + merged.min_row - source_first,
                end_row=start_row + merged.max_row - source_first,
                start_column=merged.min_col,
                end_column=merged.max_col,
            )
    for row in range(start_row, start_row + 8):
        for column in range(1, 9):
            cell = ws.cell(row, column)
            cell.alignment = copy(cell.alignment)
            cell.alignment = Alignment(
                horizontal="left", vertical=cell.alignment.vertical,
                wrap_text=cell.alignment.wrap_text, indent=cell.alignment.indent,
            )


def extract_stamp(template_path: Path, working: Path) -> Path:
    working.mkdir(parents=True, exist_ok=True)
    template_copy = working / "template-copy.xlsx"
    shutil.copyfile(template_path, template_copy)
    with zipfile.ZipFile(template_copy) as archive:
        media = next((name for name in archive.namelist() if name.startswith("xl/media/")), None)
        if not media:
            raise ValueError("No stamp image was found in template 229 WOFENG.xlsx.")
        target = working / Path(media).name
        with archive.open(media) as source, target.open("wb") as destination:
            shutil.copyfileobj(source, destination)
    return target


def pixels_for_column(ws, column: int) -> int:
    width = ws.column_dimensions[get_column_letter(column)].width
    return int((width if width is not None else 8.43) * 7 + 5)


def pixels_for_row(ws, row: int) -> int:
    height = ws.row_dimensions[row].height
    return int((height if height is not None else 15) * 4 / 3)


def add_stamp(ws, stamp_path: Path, total_row: int):
    top_row = max(20, total_row + 1)
    if top_row > 29:
        raise ValueError("No room for the stamp in D20:G29.")
    stamp = ExcelImage(stamp_path)
    available_width = sum(pixels_for_column(ws, column) for column in range(4, 8))
    available_height = sum(pixels_for_row(ws, row) for row in range(top_row, 30))
    scale = min(available_width / stamp.width, available_height / stamp.height, 1)
    stamp.width = max(1, int(stamp.width * scale))
    stamp.height = max(1, int(stamp.height * scale))
    ws.add_image(stamp, f"D{top_row}")


def prepare_file(source: Path, template: Path, destination: Path, working: Path):
    stamp_path = extract_stamp(template, working)
    shutil.copyfile(source, destination)
    workbook = load_workbook(destination)
    if "INV" not in workbook.sheetnames:
        raise ValueError(f"INV sheet is missing in {source.name}.")
    for name in list(workbook.sheetnames):
        if name != "INV":
            del workbook[name]
    ws = workbook["INV"]
    clear_outside_area(ws)
    total_row = None
    for row in range(1, 34):
        if any("total" in str(ws.cell(row, column).value).lower() for column in range(1, 10) if ws.cell(row, column).value is not None):
            total_row = row
            break
    if total_row is None:
        raise ValueError(f"Total row is missing in {source.name}.")
    block_start = total_row + 2
    if block_start + 7 > 33:
        raise ValueError(f"The block after Total does not fit A1:I33 in {source.name}.")
    template_book = load_workbook(template, data_only=False)
    if "INV" not in template_book.sheetnames:
        raise ValueError("INV sheet is missing in the payment-invoice template.")
    copy_block(template_book["INV"], ws, block_start)
    template_book.close()
    ws.column_dimensions["A"].width = 6
    beneficiary = find_text(ws, "Beneficiary Address")
    if beneficiary is None:
        raise ValueError(f"Beneficiary Address row is missing in {source.name}.")
    beneficiary_value = beneficiary.value
    unmerge_intersecting(ws, beneficiary.row, beneficiary.row, 1, 8)
    ws.merge_cells(start_row=beneficiary.row, end_row=beneficiary.row, start_column=1, end_column=8)
    beneficiary_cell = ws.cell(beneficiary.row, 1)
    beneficiary_cell.value = beneficiary_value
    beneficiary_cell.alignment = copy(beneficiary_cell.alignment)
    beneficiary_cell.alignment = Alignment(horizontal="left", vertical=beneficiary_cell.alignment.vertical, wrap_text=beneficiary_cell.alignment.wrap_text)
    ws.row_dimensions[beneficiary.row].height = 29.25
    seller = find_text(ws, "SELLER OTHER TERMS")
    if seller is None or seller.row + 2 > 33:
        raise ValueError(f"SELLER OTHER TERMS cell cannot be merged in {source.name}.")
    seller_value = seller.value
    unmerge_intersecting(ws, seller.row, seller.row + 2, seller.column, seller.column)
    ws.merge_cells(start_row=seller.row, end_row=seller.row + 2, start_column=seller.column, end_column=seller.column)
    seller_cell = ws.cell(seller.row, seller.column)
    seller_cell.value = seller_value
    seller_cell.alignment = Alignment(horizontal="left", vertical=seller_cell.alignment.vertical, wrap_text=seller_cell.alignment.wrap_text)
    add_stamp(ws, stamp_path, total_row)
    workbook.save(destination)
    workbook.close()


def validate(destination: Path):
    workbook = load_workbook(destination, data_only=False)
    try:
        if workbook.sheetnames != ["INV"]:
            raise ValueError("Validation failed: only INV must remain.")
        ws = workbook["INV"]
        for cell in ws._cells.values():
            if (cell.row > 33 or cell.column > 9) and cell.value is not None:
                raise ValueError("Validation failed: data remains outside A1:I33.")
        if len(ws._images) != 1:
            raise ValueError("Validation failed: exactly one stamp image is required.")
        if ws.column_dimensions["A"].width != 6:
            raise ValueError("Validation failed: column A width is incorrect.")
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--job-dir", required=True)
    parser.add_argument("--template", required=True)
    args = parser.parse_args()
    job_dir = Path(args.job_dir)
    template = Path(args.template)
    output_dir = job_dir / "output"
    output_dir.mkdir(exist_ok=True)
    inputs = list((job_dir / "input").glob("*.xlsx"))
    if not inputs:
        raise ValueError("Task 6 needs source .xlsx files.")
    output_files, notes = [], []
    for index, source in enumerate(inputs, start=1):
        update_progress(job_dir, 5 + round((index - 1) / len(inputs) * 90), f'Обрабатываю файл {index} из {len(inputs)}: {source.name}')
        name = output_name(source.stem)
        destination = output_dir / name
        if destination.exists():
            raise ValueError(f"Output file already exists: {name}")
        prepare_file(source, template, destination, job_dir / "working")
        validate(destination)
        output_files.append(name)
        notes.append(f"{name}: INV prepared from template 229 WOFENG.")
        update_progress(job_dir, 5 + round(index / len(inputs) * 90), f'Обработано файлов: {index} из {len(inputs)}')
    print(json.dumps({"outputFiles": output_files, "notes": notes}, ensure_ascii=True))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)

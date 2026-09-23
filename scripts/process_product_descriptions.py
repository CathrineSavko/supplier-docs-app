import argparse
import copy
import hashlib
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from docx import Document

from process_commercial_offer import normalize, read_pckb, supplier_key, update_progress


ARTICLE = re.compile(r'^[A-Z]{2,10}-\d{3,}$', re.I)


def article_key(value):
    return re.sub(r'\s+', '', normalize(value)).upper()


def hash_file(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def find_template(templates_dir, supplier):
    supplier = supplier.casefold()
    matches = []
    for path in Path(templates_dir).rglob('*.docx'):
        name = path.name.casefold()
        if supplier in name and ('общее' in name or 'general' in name):
            matches.append(path)
    if len(matches) != 1:
        raise ValueError('Description template is missing or ambiguous.')
    return matches[0]


def row_article(row):
    if len(row.cells) < 2:
        return ''
    return article_key(row.cells[1].text)


def image_count(row):
    return sum(
        1 for element in row._tr.iter()
        if element.tag.rsplit('}', 1)[-1] in {'blip', 'imagedata'}
    )


def anchors_in_document(document):
    return sum(1 for element in document.element.body.iter() if element.tag.rsplit('}', 1)[-1] == 'anchor')


def row_anchor_count(row):
    return sum(1 for element in row._tr.iter() if element.tag.rsplit('}', 1)[-1] == 'anchor')


def anchors_outside_product_table(document, table):
    table_element = table._tbl
    total = 0
    for anchor in (element for element in document.element.body.iter() if element.tag.rsplit('}', 1)[-1] == 'anchor'):
        parent = anchor.getparent()
        while parent is not None and parent is not table_element:
            parent = parent.getparent()
        if parent is None:
            total += 1
    return total


def set_cell_text(cell, value):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = str(value)
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(str(value))
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def set_document_date(document, date_value):
    pattern = re.compile(r'(Исх\.\s*б/н\s*от\s*)\d{2}\.\d{2}\.\d{4}')
    replacement = date_value.strftime('%d.%m.%Y')
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if pattern.search(run.text):
                run.text = pattern.sub(r'\g<1>' + replacement, run.text)
                return
        if pattern.search(paragraph.text):
            if paragraph.runs:
                paragraph.runs[0].text = pattern.sub(r'\g<1>' + replacement, paragraph.text)
                for run in paragraph.runs[1:]:
                    run._element.getparent().remove(run._element)
                return
    raise ValueError('Document date field was not found in the description template.')


def catalogue_map(table):
    mapping = {}
    for row in table.rows[1:]:
        article = row_article(row)
        if not ARTICLE.fullmatch(article) or image_count(row) == 0:
            continue
        mapping.setdefault(article, row)
    return mapping


def build_document(template, destination, rows, document_date):
    document = Document(template)
    if len(document.tables) != 1:
        raise ValueError('The description template must contain exactly one product table.')
    table = document.tables[0]
    if not table.rows:
        raise ValueError('The description template product table is empty.')
    mapping = catalogue_map(table)
    missing = [item['article'] for item in rows if article_key(item['article']) not in mapping]
    if missing:
        raise ValueError('No exact catalogue row with photo for: ' + ', '.join(missing))

    preserved_anchor_count = anchors_outside_product_table(document, table)
    output_anchor_count = preserved_anchor_count + sum(row_anchor_count(mapping[article_key(item['article'])]) for item in rows)
    header = copy.deepcopy(table.rows[0]._tr)
    for row in list(table.rows):
        table._tbl.remove(row._tr)
    table._tbl.append(header)

    for index, item in enumerate(rows, start=1):
        source_row = mapping[article_key(item['article'])]
        table._tbl.append(copy.deepcopy(source_row._tr))
        output_row = table.rows[-1]
        set_cell_text(output_row.cells[0], index)
        set_cell_text(output_row.cells[1], item['article'])
        set_cell_text(output_row.cells[2], item['name'])

    set_document_date(document, document_date)
    if anchors_outside_product_table(document, table) != preserved_anchor_count:
        raise ValueError('Template stamps or anchored images changed unexpectedly.')
    document.save(destination)
    return preserved_anchor_count, output_anchor_count


def validate_document(destination, expected, expected_date, preserved_anchor_count, output_anchor_count):
    document = Document(destination)
    if len(document.tables) != 1:
        raise ValueError('Validation failed: product table is missing.')
    table = document.tables[0]
    rows = table.rows[1:]
    if len(rows) != len(expected):
        raise ValueError('Validation failed: product row count differs from PkCB.')
    for index, (row, item) in enumerate(zip(rows, expected), start=1):
        if normalize(row.cells[0].text) != str(index):
            raise ValueError(f'Validation failed: wrong sequence number in product row {index}.')
        if article_key(row.cells[1].text) != article_key(item['article']):
            raise ValueError(f'Validation failed: wrong article in product row {index}.')
        if normalize(row.cells[2].text) != normalize(item['name']):
            raise ValueError(f'Validation failed: wrong product name in product row {index}.')
        description = ' '.join(normalize(cell.text) for cell in row.cells[3:-1] or row.cells[3:])
        if not description:
            raise ValueError(f'Validation failed: product description is missing in row {index}.')
        if image_count(row) == 0:
            raise ValueError(f'Validation failed: product photo is missing in row {index}.')
    expected_date_text = expected_date.strftime('%d.%m.%Y')
    if not any(expected_date_text in paragraph.text and 'Исх.' in paragraph.text for paragraph in document.paragraphs):
        raise ValueError('Validation failed: document date is incorrect.')
    if anchors_outside_product_table(document, table) != preserved_anchor_count or anchors_in_document(document) != output_anchor_count:
        raise ValueError('Validation failed: template stamps or anchored images are missing.')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--job-dir', required=True)
    parser.add_argument('--templates-dir', required=True)
    args = parser.parse_args()
    job_dir = Path(args.job_dir)
    request = json.loads((job_dir / 'request.json').read_text(encoding='utf-8'))
    parameters = request.get('parameters', {})
    supplier = supplier_key(parameters.get('supplier'))
    pc_number = normalize(parameters.get('pcNumber'))
    containers = normalize(parameters.get('containers'))
    if not pc_number or not containers:
        raise ValueError('PkCB number and container numbers are required.')
    try:
        document_date = datetime.strptime(parameters.get('documentDate', ''), '%Y-%m-%d').date()
    except ValueError as error:
        raise ValueError('Document date is required.') from error
    sources = list((job_dir / 'input').glob('*.xls'))
    if len(sources) != 1:
        raise ValueError('Task 4 needs exactly one PkCB .xls file.')

    update_progress(job_dir, 8, 'Reading PkCB...')
    pckb_rows = read_pckb(sources[0])
    if any(not item.get('name') for item in pckb_rows):
        raise ValueError('One or more PkCB product names are missing.')
    update_progress(job_dir, 25, f'Found product positions: {len(pckb_rows)}')
    template = find_template(args.templates_dir, supplier)
    template_hash = hash_file(template)
    output_dir = job_dir / 'output'
    output_dir.mkdir(exist_ok=True)
    name = f'Описание {supplier.title()} {containers}.docx'
    destination = output_dir / name
    update_progress(job_dir, 45, 'Copying descriptions and product photos...')
    preserved_anchor_count, output_anchor_count = build_document(template, destination, pckb_rows, document_date)
    update_progress(job_dir, 88, 'Checking the finished Word document...')
    validate_document(destination, pckb_rows, document_date, preserved_anchor_count, output_anchor_count)
    if hash_file(template) != template_hash:
        raise ValueError('The retained description template changed unexpectedly.')
    print(json.dumps({
        'outputFiles': [name],
        'notes': [f'Product descriptions: {len(pckb_rows)} positions; exact article-photo matching verified.'],
    }, ensure_ascii=True))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)

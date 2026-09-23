import argparse
import json
import re
import shutil
import struct
import sys
import warnings
from copy import copy
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels

warnings.filterwarnings('ignore', message='wmf image format is not supported so the image is being dropped')


SKU = re.compile(r'\b[A-Z]{2,10}-\d{3,}\b', re.I)
END_OF_CHAIN = 0xFFFFFFFE
FREE_SECTOR = 0xFFFFFFFF


def update_progress(job_dir, percent, stage):
    request = job_dir / 'request.json'
    try:
        data = json.loads(request.read_text(encoding='utf-8'))
        data['progress'] = percent
        data['stage'] = stage
        request.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    except Exception:
        pass


def normalize(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def supplier_key(value):
    key = normalize(value).casefold()
    if key not in {'wofeng', 'gumi'}:
        raise ValueError('Supplier must be Wofeng or GUMI.')
    return key


def read_u16(data, offset):
    return struct.unpack_from('<H', data, offset)[0]


def read_u32(data, offset):
    return struct.unpack_from('<I', data, offset)[0]


class CompoundFile:
    """Small read-only reader for the BIFF8 .xls files used as PkCB inputs."""

    def __init__(self, path):
        self.data = Path(path).read_bytes()
        if self.data[:8] != b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
            raise ValueError('PkCB source must be an Excel .xls file.')
        self.sector_size = 1 << read_u16(self.data, 30)
        self.mini_sector_size = 1 << read_u16(self.data, 32)
        self.first_dir_sector = read_u32(self.data, 48)
        self.mini_cutoff = read_u32(self.data, 56)
        self.first_mini_fat = read_u32(self.data, 60)
        self.num_mini_fat = read_u32(self.data, 64)
        self.first_difat = read_u32(self.data, 68)
        self.num_difat = read_u32(self.data, 72)
        self.fat = self._read_fat()
        self.entries, self.root = self._read_directory()
        self.mini_stream = self._stream(self.root['start'], self.root['size'])
        self.mini_fat = self._read_mini_fat()

    def _sector(self, sector_id):
        start = (sector_id + 1) * self.sector_size
        return self.data[start:start + self.sector_size]

    def _read_fat(self):
        difat = [read_u32(self.data, 76 + 4 * index) for index in range(109)]
        next_sector = self.first_difat
        for _ in range(self.num_difat):
            block = self._sector(next_sector)
            difat.extend(read_u32(block, 4 * index) for index in range(self.sector_size // 4 - 1))
            next_sector = read_u32(block, self.sector_size - 4)
        result = []
        for sector_id in difat:
            if sector_id in {FREE_SECTOR, END_OF_CHAIN}:
                continue
            block = self._sector(sector_id)
            result.extend(read_u32(block, 4 * index) for index in range(self.sector_size // 4))
        return result

    def _chain(self, start, table=None):
        table = self.fat if table is None else table
        seen, sectors, current = set(), [], start
        while current not in {END_OF_CHAIN, FREE_SECTOR} and current not in seen:
            if current >= len(table):
                raise ValueError('Invalid Excel .xls sector chain.')
            seen.add(current)
            sectors.append(current)
            current = table[current]
        return sectors

    def _stream(self, start, size):
        return b''.join(self._sector(item) for item in self._chain(start))[:size]

    def _read_directory(self):
        entries, root = {}, None
        raw = self._stream(self.first_dir_sector, len(self.data))
        for offset in range(0, len(raw) - 128 + 1, 128):
            item = raw[offset:offset + 128]
            length = read_u16(item, 64)
            kind = item[66]
            if not length or kind not in {2, 5}:
                continue
            name = item[:max(0, length - 2)].decode('utf-16le', errors='ignore')
            entry = {'name': name, 'kind': kind, 'start': read_u32(item, 116), 'size': struct.unpack_from('<Q', item, 120)[0]}
            if kind == 5:
                root = entry
            else:
                entries[name.casefold()] = entry
        if root is None:
            raise ValueError('Excel .xls root stream is missing.')
        return entries, root

    def _read_mini_fat(self):
        if self.num_mini_fat == 0 or self.first_mini_fat in {END_OF_CHAIN, FREE_SECTOR}:
            return []
        raw = self._stream(self.first_mini_fat, self.num_mini_fat * self.sector_size)
        return [read_u32(raw, offset) for offset in range(0, len(raw), 4)]

    def stream(self, name):
        entry = self.entries.get(name.casefold())
        if entry is None:
            raise ValueError('Workbook stream is missing from Excel .xls.')
        if entry['size'] >= self.mini_cutoff:
            return self._stream(entry['start'], entry['size'])
        chunks = []
        for sector in self._chain(entry['start'], self.mini_fat):
            start = sector * self.mini_sector_size
            chunks.append(self.mini_stream[start:start + self.mini_sector_size])
        return b''.join(chunks)[:entry['size']]


def biff_records(data, start=0):
    position = start
    while position + 4 <= len(data):
        record_id, size = struct.unpack_from('<HH', data, position)
        position += 4
        payload = data[position:position + size]
        position += size
        yield record_id, payload


def decode_biff_text(raw, compressed=True):
    if compressed:
        for encoding in ('cp1251', 'latin1'):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                pass
    return raw.decode('utf-16le', errors='replace')


def read_sst(records):
    payloads = []
    index = 0
    while index < len(records):
        record_id, payload = records[index]
        if record_id == 0x00FC:
            payloads.append(payload)
            index += 1
            while index < len(records) and records[index][0] == 0x003C:
                payloads.append(records[index][1])
                index += 1
            break
        index += 1
    if not payloads:
        return []
    data = b''.join(payloads)
    if len(data) < 8:
        return []
    count, position, values = read_u32(data, 4), 8, []
    for _ in range(count):
        if position + 3 > len(data):
            break
        chars = read_u16(data, position)
        options = data[position + 2]
        position += 3
        rich_count = read_u16(data, position) if options & 0x08 else 0
        if options & 0x08:
            position += 2
        phonetic_size = read_u32(data, position) if options & 0x04 else 0
        if options & 0x04:
            position += 4
        width = 2 if options & 0x01 else 1
        size = chars * width
        if position + size > len(data):
            break
        values.append(decode_biff_text(data[position:position + size], compressed=width == 1))
        position += size + rich_count * 4 + phonetic_size
    return values


def rk_value(value):
    integer = bool(value & 0x02)
    divided = bool(value & 0x01)
    if integer:
        number = value >> 2
        if number & (1 << 29):
            number -= 1 << 30
    else:
        number = struct.unpack('<d', b'\0\0\0\0' + struct.pack('<I', value & 0xFFFFFFFC))[0]
    return number / 100 if divided else number


def sheet_cells(data, offset, shared_strings):
    cells = {}
    for record_id, payload in biff_records(data, offset):
        if record_id == 0x000A:
            break
        if record_id == 0x0204 and len(payload) >= 8:
            row, col, chars = read_u16(payload, 0), read_u16(payload, 2), read_u16(payload, 6)
            cells[(row, col)] = decode_biff_text(payload[8:8 + chars], compressed=True)
        elif record_id in {0x0203, 0x027E, 0x00FD, 0x0006, 0x0205} and len(payload) >= 8:
            row, col = read_u16(payload, 0), read_u16(payload, 2)
            if record_id == 0x0203:
                value = struct.unpack_from('<d', payload, 6)[0]
            elif record_id == 0x027E:
                value = rk_value(read_u32(payload, 6))
            elif record_id == 0x00FD:
                string_index = read_u32(payload, 6)
                value = shared_strings[string_index] if string_index < len(shared_strings) else ''
            elif record_id == 0x0205:
                value = bool(payload[6])
            else:
                value = struct.unpack_from('<d', payload, 6)[0]
            cells[(row, col)] = value
        elif record_id == 0x00BD and len(payload) >= 10:
            row, first_col = read_u16(payload, 0), read_u16(payload, 2)
            last_col = read_u16(payload, len(payload) - 2)
            position = 4
            for col in range(first_col, last_col + 1):
                if position + 6 > len(payload) - 2:
                    break
                cells[(row, col)] = rk_value(read_u32(payload, position + 2))
                position += 6
    return cells


def read_pckb(source):
    container = CompoundFile(source)
    try:
        stream = container.stream('Workbook')
    except ValueError:
        stream = container.stream('Book')
    records = list(biff_records(stream))
    strings, sheets = read_sst(records), []
    for record_id, payload in records:
        if record_id != 0x0085 or len(payload) < 8:
            continue
        offset = read_u32(payload, 0)
        chars, flags = payload[6], payload[7]
        width = 2 if flags & 0x01 else 1
        name = decode_biff_text(payload[8:8 + chars * width], compressed=width == 1)
        sheets.append((name, offset))
    rows = []
    for _, offset in sheets:
        cells = sheet_cells(stream, offset, strings)
        by_row = {}
        for (row, col), value in cells.items():
            by_row.setdefault(row, {})[col] = value
        rows.extend(by_row.values())
    positions = []
    for row in rows:
        item_columns = [(column, normalize(value).upper()) for column, value in row.items() if isinstance(value, str) and SKU.fullmatch(normalize(value))]
        if not item_columns:
            continue
        column, article = item_columns[0]
        numbers = [(key, float(value)) for key, value in sorted(row.items()) if key > column and isinstance(value, (int, float)) and not isinstance(value, bool)]
        result = None
        for first in range(len(numbers)):
            for second in range(first + 1, len(numbers)):
                for third in range(second + 1, len(numbers)):
                    quantity, price, amount = numbers[first][1], numbers[second][1], numbers[third][1]
                    if quantity > 0 and price >= 0 and abs(quantity * price - amount) < 0.02:
                        result = quantity, price, amount
                        break
                if result:
                    break
            if result:
                break
        if result:
            positions.append({
                'article': article,
                'name': normalize(row.get(column + 1, '')),
                'quantity': result[0],
                'price': result[1],
                'amount': result[2],
            })
    if not positions:
        raise ValueError('No product rows were found in the PkCB .xls file.')
    return positions


def find_template(templates_dir, supplier, marker):
    supplier = supplier.casefold()
    marker = marker.casefold()
    matches = []
    for path in Path(templates_dir).rglob('*.xlsx'):
        name = path.name.casefold()
        if supplier in name and marker in name and 'commercial offer' in name:
            matches.append(path)
    if len(matches) != 1:
        label = 'empty template' if marker == 'пустой' else 'catalog template'
        raise ValueError(f'Missing or ambiguous {label} for {supplier.title()}. Add it in Templates.')
    return matches[0]


def find_cell(ws, predicate):
    for row in ws.iter_rows():
        for cell in row:
            if predicate(normalize(cell.value)):
                return cell
    return None


def cell_image_row(image):
    anchor = image.anchor
    marker = getattr(anchor, '_from', None)
    return marker.row + 1 if marker else None


def catalog_items(path):
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        description_column = find_cell(sheet, lambda value: value.casefold() == 'designation')
        if description_column is None:
            raise ValueError('Designation column was not found in the catalog.')
        images = {cell_image_row(image): image._data() for image in sheet._images if cell_image_row(image)}
        result = {}
        for row in range(description_column.row + 1, sheet.max_row + 1):
            description = normalize(sheet.cell(row, description_column.column).value)
            for article in SKU.findall(description.upper()):
                result[article] = {'designation': description, 'image': images.get(row)}
        return result
    finally:
        workbook.close()


def copy_row_style(ws, source_row, target_row):
    for column in range(1, ws.max_column + 1):
        source, target = ws.cell(source_row, column), ws.cell(target_row, column)
        target._style = copy(source._style)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
        target.alignment = copy(source.alignment)
        target.fill = copy(source.fill)
        target.font = copy(source.font)
        target.border = copy(source.border)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def shift_images(ws, first_row, delta):
    for image in ws._images:
        anchor = image.anchor
        marker = getattr(anchor, '_from', None)
        if marker and marker.row + 1 >= first_row:
            marker.row += delta
            last = getattr(anchor, 'to', None)
            if last:
                last.row += delta


def fit_image(image, ws, row, column):
    letter = get_column_letter(column)
    width = ws.column_dimensions[letter].width or 12
    row_height = ws.row_dimensions[row].height
    if row_height is None:
        row_height = 45
        ws.row_dimensions[row].height = row_height
    cell_width = max(32, int((width + 0.75) * 7))
    cell_height = max(32, points_to_pixels(row_height))
    max_width = max(24, cell_width - 20)
    max_height = max(24, cell_height - 18)
    scale = min(max_width / image.width, max_height / image.height, 1)
    image.width = max(1, int(image.width * scale))
    image.height = max(1, int(image.height * scale))
    image.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=column - 1,
            colOff=pixels_to_EMU((cell_width - image.width) // 2),
            row=row - 1,
            rowOff=pixels_to_EMU((cell_height - image.height) // 2),
        ),
        ext=XDRPositiveSize2D(cx=pixels_to_EMU(image.width), cy=pixels_to_EMU(image.height)),
    )


def clone_image(image_data):
    return ExcelImage(BytesIO(image_data))


def english_date(value):
    return f'{value.day:02d} {value.strftime("%B")} {value.year}'


def money(value):
    return f'{value:,.2f}'.replace(',', ' ').replace('.', ',')


def prepare_table(ws, count):
    designation = find_cell(ws, lambda value: value.casefold() == 'designation')
    pictures = find_cell(ws, lambda value: value.casefold() == 'pictures')
    if designation is None or pictures is None:
        raise ValueError('Commercial offer table headers were not found.')
    start = designation.row + 2
    total_label = find_cell(ws, lambda value: value.casefold().startswith('total'))
    if total_label is None or total_label.row <= start:
        raise ValueError('Commercial offer Total row was not found.')
    total_row, capacity = total_label.row, total_label.row - start
    if count > capacity:
        delta = count - capacity
        shift_images(ws, total_row, delta)
        ws.insert_rows(total_row, delta)
        for row in range(total_row, total_row + delta):
            copy_row_style(ws, total_row - 1, row)
        total_row += delta
    elif count < capacity:
        delta = capacity - count
        shift_images(ws, start + count, -delta)
        ws.delete_rows(start + count, delta)
        total_row -= delta
    return {'start': start, 'end': start + count - 1, 'total': total_row, 'designation_col': designation.column, 'pictures_col': pictures.column}


def format_product_table(ws, table):
    header = table['start'] - 2
    rows = [header, *range(table['start'], table['end'] + 1), table['total']]
    for row in rows:
        for column in range(1, 8):
            cell = ws.cell(row, column)
            font = copy(cell.font)
            font.name = 'Arial'
            font.sz = 9
            cell.font = font
            alignment = copy(cell.alignment)
            alignment.horizontal = 'center'
            alignment.vertical = 'center'
            alignment.wrap_text = True
            cell.alignment = alignment
    edge = Side(style='thin', color='000000')
    for row in rows:
        ws.cell(row, table['designation_col']).border = Border(left=edge, right=edge, top=edge, bottom=edge)


def update_footer(ws, table, document_date, total):
    date_label = find_cell(ws, lambda value: value.casefold().rstrip(':') == 'date')
    if date_label:
        ws.cell(date_label.row + 1, date_label.column).value = english_date(document_date)
    footer = find_cell(ws, lambda value: 'total amount:' in value.casefold())
    if footer is None:
        raise ValueError('Commercial offer footer is missing TOTAL AMOUNT.')
    footer.value = f'1.TOTAL AMOUNT: RMB {money(total)}'
    ws.cell(footer.row + 1, footer.column).value = '2.100% in 25 calendar days after the shipment'
    ws.cell(footer.row + 2, footer.column).value = '3.DELIVERY TIME: 90 DAYS'
    valid_until = document_date + timedelta(days=90)
    ws.cell(footer.row + 3, footer.column).value = f'4.VALIDITY: THIS COMMERCIAL OFFERTA IS VALID FROM {document_date:%d.%m.%Y} to {valid_until:%d.%m.%Y}'
    ws.cell(table['total'], 7).value = f'=SUM(G{table["start"]}:G{table["end"]})'


def build_offer(blank_template, catalog_template, destination, supplier, document_date, pckb_rows):
    catalog = catalog_items(catalog_template)
    missing = [item['article'] for item in pckb_rows if item['article'] not in catalog or not catalog[item['article']]['image']]
    if missing:
        raise ValueError('No matching catalog photo was found for: ' + ', '.join(missing))
    shutil.copyfile(blank_template, destination)
    workbook = load_workbook(destination, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        seals = len(sheet._images)
        table = prepare_table(sheet, len(pckb_rows))
        for row in range(table['start'], table['end'] + 1):
            for column in range(1, sheet.max_column + 1):
                sheet.cell(row, column).value = None
        total = 0
        for index, item in enumerate(pckb_rows, start=1):
            row, match = table['start'] + index - 1, catalog[item['article']]
            designation = match['designation']
            if item['article'] not in designation.upper():
                designation = f'{designation} ({item["article"]})'
            sheet.cell(row, 1).value = index
            sheet.cell(row, table['designation_col']).value = designation
            sheet.cell(row, 4).value = item['quantity']
            sheet.cell(row, 5).value = 'pcs'
            sheet.cell(row, 6).value = item['price']
            sheet.cell(row, 7).value = item['amount']
            picture = clone_image(match['image'])
            fit_image(picture, sheet, row, table['pictures_col'])
            sheet.add_image(picture)
            total += item['amount']
        format_product_table(sheet, table)
        update_footer(sheet, table, document_date, total)
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(destination)
    finally:
        workbook.close()
    return {'count': len(pckb_rows), 'total': total, 'seals': seals}


def validate_offer(destination, expected, document_date):
    workbook = load_workbook(destination, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        table = prepare_table(sheet, expected['count'])
        rows = []
        for row in range(table['start'], table['end'] + 1):
            values = [sheet.cell(row, column).value for column in (1, table['designation_col'], 4, 6, 7)]
            if any(value is None for value in values):
                raise ValueError(f'Validation failed: incomplete product row {row}.')
            rows.append(values)
        if len(rows) != expected['count']:
            raise ValueError('Validation failed: product row count differs from PkCB.')
        if abs(sum(float(row[4]) for row in rows) - expected['total']) > 0.02:
            raise ValueError('Validation failed: product total differs from PkCB.')
        if len(sheet._images) < expected['seals'] + expected['count']:
            raise ValueError('Validation failed: one or more product photos or seals are missing.')
        photos = []
        for image in sheet._images:
            anchor = image.anchor
            marker = getattr(anchor, '_from', None)
            extent = getattr(anchor, 'ext', None)
            if marker and extent and marker.col == table['pictures_col'] - 1 and table['start'] - 1 <= marker.row <= table['end'] - 1:
                photos.append(image)
                width = sheet.column_dimensions[get_column_letter(table['pictures_col'])].width or 12
                cell_width = max(32, int((width + 0.75) * 7))
                cell_height = max(32, points_to_pixels(sheet.row_dimensions[marker.row + 1].height or 45))
                image_width, image_height = extent.cx / 9525, extent.cy / 9525
                offset_x, offset_y = marker.colOff / 9525, marker.rowOff / 9525
                if offset_x < 0 or offset_y < 0 or offset_x + image_width > cell_width or offset_y + image_height > cell_height:
                    raise ValueError('Validation failed: a product photo leaves its PICTURES cell.')
        if len(photos) != expected['count']:
            raise ValueError('Validation failed: product photos do not match the product row count.')
        rows_to_check = [table['start'] - 2, *range(table['start'], table['end'] + 1), table['total']]
        for row in rows_to_check:
            for column in range(1, 8):
                cell = sheet.cell(row, column)
                if cell.font.name != 'Arial' or cell.font.sz != 9 or cell.alignment.horizontal != 'center' or cell.alignment.vertical != 'center':
                    raise ValueError(f'Validation failed: table format is incorrect at {cell.coordinate}.')
            border = sheet.cell(row, table['designation_col']).border
            if not all(side.style == 'thin' and side.color and side.color.rgb in {'00000000', '000000'} for side in (border.left, border.right, border.top, border.bottom)):
                raise ValueError(f'Validation failed: designation border is missing at C{row}.')
        date_label = find_cell(sheet, lambda value: value.casefold().rstrip(':') == 'date')
        if date_label and sheet.cell(date_label.row + 1, date_label.column).value != english_date(document_date):
            raise ValueError('Validation failed: document date is incorrect.')
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--job-dir', required=True)
    parser.add_argument('--templates-dir', required=True)
    args = parser.parse_args()
    job_dir = Path(args.job_dir)
    request = json.loads((job_dir / 'request.json').read_text(encoding='utf-8'))
    parameters = request.get('parameters', {})
    supplier = supplier_key(parameters.get('supplier'))
    pc_number = normalize(parameters.get('pcNumber'))
    if not pc_number:
        raise ValueError('PkCB number is required.')
    try:
        document_date = datetime.strptime(parameters.get('documentDate', ''), '%Y-%m-%d').date()
    except ValueError as error:
        raise ValueError('Document date is required.') from error
    sources = list((job_dir / 'input').glob('*.xls')) + list((job_dir / 'input').glob('*.xlsx'))
    if len(sources) != 1:
        raise ValueError('Task 3 needs exactly one PkCB .xls file.')
    update_progress(job_dir, 8, 'Читаю ПкЦБ…')
    pckb_rows = read_pckb(sources[0])
    update_progress(job_dir, 25, f'Найдено товарных позиций: {len(pckb_rows)}')
    blank = find_template(args.templates_dir, supplier, 'пустой')
    catalog = find_template(args.templates_dir, supplier, 'всё')
    output_dir = job_dir / 'output'
    output_dir.mkdir(exist_ok=True)
    name = f'ПкЦБ-{pc_number} Commercial offer {supplier.title()}.xlsx'
    destination = output_dir / name
    update_progress(job_dir, 45, 'Переношу товары и фотографии в шаблон…')
    details = build_offer(blank, catalog, destination, supplier, document_date, pckb_rows)
    update_progress(job_dir, 88, 'Проверяю готовый Commercial offer…')
    validate_offer(destination, details, document_date)
    print(json.dumps({'outputFiles': [name], 'notes': [f'Commercial offer: {details["count"]} positions, total RMB {money(details["total"])}.']}, ensure_ascii=True))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)

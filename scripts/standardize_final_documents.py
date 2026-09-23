import argparse
import json
import re
import shutil
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, PatternFill


def update_progress(job_dir, percent, stage):
    request = job_dir / 'request.json'
    try:
        data = json.loads(request.read_text(encoding='utf-8'))
        data['progress'] = percent
        data['stage'] = stage
        request.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    except Exception:
        pass


def normalized_name(stem):
    match = re.search(r'CTI\d{2}(\d{3})\s*BY', stem, flags=re.I)
    if not match:
        match = re.search(r'(\d{3})(?:\D|$)', stem)
    if not match:
        raise ValueError(f'Container number was not recognised in {stem}.')
    return f'{match.group(1)} WOFENG.xlsx'


def find_table(ws):
    item, description = None, None
    for row in ws.iter_rows():
        for cell in row:
            value = re.sub(r'\s+', ' ', str(cell.value or '')).strip().casefold()
            if 'item' in value and 'no' in value:
                item = cell
            if 'description' in value:
                description = cell
    if not item or not description:
        return None
    header_row = max(item.row, description.row)
    total_row = None
    for row in range(header_row + 1, ws.max_row + 1):
        if any(re.sub(r'\s+', ' ', str(ws.cell(row, column).value or '')).strip().casefold().startswith('total') for column in range(1, ws.max_column + 1)):
            total_row = row
            break
    if total_row is None:
        return None
    return header_row, total_row


def snapshot(ws):
    return {
        'view': ws.sheet_view.view,
        'zoom': ws.sheet_view.zoomScale,
        'zoom_normal': ws.sheet_view.zoomScaleNormal,
        'print_area': str(ws.print_area),
        'images': len(ws._images),
    }


def apply_sheet_format(ws):
    table = find_table(ws)
    if not table:
        raise ValueError(f'Item table was not recognised on {ws.title}.')
    first_row, last_row = table
    for row in ws.iter_rows():
        for cell in row:
            if cell.fill.fill_type is not None:
                cell.fill = PatternFill()
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            font = copy(cell.font)
            cell.font = Font(
                name='Arial Cyr', sz=8, bold=True, italic=font.italic,
                vertAlign=font.vertAlign, underline=font.underline,
                strike=font.strike, color=copy(font.color), charset=font.charset,
                family=font.family, scheme=font.scheme, outline=font.outline,
                shadow=font.shadow, condense=font.condense, extend=font.extend,
            )


def validate(destination, before):
    workbook = load_workbook(destination, data_only=False)
    try:
        for name in ('INV', 'PL'):
            if name not in workbook.sheetnames:
                raise ValueError(f'Validation failed: {name} is missing.')
            ws = workbook[name]
            after = snapshot(ws)
            for key in ('view', 'zoom', 'zoom_normal', 'print_area', 'images'):
                if after[key] != before[name][key]:
                    raise ValueError(f'Validation failed: {name} {key} changed.')
            table = find_table(ws)
            if not table:
                raise ValueError(f'Validation failed: item table missing on {name}.')
            for row in ws.iter_rows(min_row=table[0], max_row=table[1], min_col=1, max_col=ws.max_column):
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    if cell.font.name != 'Arial Cyr' or cell.font.sz != 8 or not cell.font.bold:
                        raise ValueError(f'Validation failed: font is incorrect at {name}!{cell.coordinate}.')
            for row in ws.iter_rows():
                for cell in row:
                    if cell.fill.fill_type is not None:
                        raise ValueError(f'Validation failed: fill remains at {name}!{cell.coordinate}.')
    finally:
        workbook.close()


def process(source, destination):
    shutil.copyfile(source, destination)
    workbook = load_workbook(destination, data_only=False)
    try:
        before = {}
        for name in ('INV', 'PL'):
            if name not in workbook.sheetnames:
                raise ValueError(f'{name} sheet is missing in {source.name}.')
            before[name] = snapshot(workbook[name])
            apply_sheet_format(workbook[name])
        workbook.save(destination)
    finally:
        workbook.close()
    validate(destination, before)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--job-dir', required=True)
    args = parser.parse_args()
    job_dir = Path(args.job_dir)
    sources = list((job_dir / 'input').glob('*.xlsx'))
    if not sources:
        raise ValueError('Task 2 needs source .xlsx files.')
    output_dir = job_dir / 'output'
    output_dir.mkdir(exist_ok=True)
    output_files, notes = [], []
    for index, source in enumerate(sources, start=1):
        update_progress(job_dir, 5 + round((index - 1) / len(sources) * 90), f'Оформляю файл {index} из {len(sources)}: {source.name}')
        name = normalized_name(source.stem)
        destination = output_dir / name
        if destination.exists():
            raise ValueError(f'Output file already exists: {name}')
        process(source, destination)
        output_files.append(name)
        notes.append(f'{name}: INV and PL formatted; fills removed.')
        update_progress(job_dir, 5 + round(index / len(sources) * 90), f'Оформлено файлов: {index} из {len(sources)}')
    print(json.dumps({'outputFiles': output_files, 'notes': notes}, ensure_ascii=True))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
